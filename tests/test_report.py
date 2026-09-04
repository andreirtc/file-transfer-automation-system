"""
Unit tests for ReportService, corporate checklist report generation,
dynamic row expansion, automated integrity checking, and SQLite date queries.
"""

import os
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from core.models import FileStatus, TransferJob, TransferRecord
from services.configuration_service import ConfigurationService
from services.database_service import DatabaseService
from services.report_service import ReportService


def test_default_checklist_systems(config):
    """Verify default checklist systems and sign-off defaults."""
    systems = config.checklist_systems
    assert len(systems) == 8
    job_names = [s["job_name"] for s in systems]
    assert "TFS42PROD" in job_names
    assert "CSE" in job_names
    assert "TFA" in job_names
    assert "COGNOS" in job_names
    assert "SAP" in job_names
    assert "GOCANVAS" in job_names
    assert "DPPS" in job_names
    assert "QMS" in job_names

    assert config.report_checked_by == "Philip M. Bayudan"
    assert config.report_repository_tag == "TFSPH-PRIMARY-REPO"
    assert config.report_auto_generate is True


def test_database_get_records_by_date(test_db):
    """Verify get_records_by_date returns records filtered by transfer_completed/detected_at."""
    job = TransferJob(name="TestJob", source_folder="/src", destination_folder="/dst")
    test_db.save_job(job)

    rec_today = TransferRecord(
        job_id=job.id,
        file_name="TFS42PROD_20260904.zip",
        source_path="/src/today.zip",
        destination_path="/dst/today.zip",
        status=FileStatus.COMPLETED,
        transfer_completed=datetime(2026, 9, 4, 14, 30, 0),
        verification_passed=True,
    )
    rec_yesterday = TransferRecord(
        job_id=job.id,
        file_name="TFS42PROD_20260903.zip",
        source_path="/src/yesterday.zip",
        destination_path="/dst/yesterday.zip",
        status=FileStatus.COMPLETED,
        transfer_completed=datetime(2026, 9, 3, 10, 15, 0),
        verification_passed=True,
    )

    test_db.save_record(rec_today)
    test_db.save_record(rec_yesterday)

    today_records = test_db.get_records_by_date(date(2026, 9, 4))
    assert len(today_records) == 1
    assert today_records[0].file_name == "TFS42PROD_20260904.zip"

    yesterday_records = test_db.get_records_by_date("2026-09-03")
    assert len(yesterday_records) == 1
    assert yesterday_records[0].file_name == "TFS42PROD_20260903.zip"


def test_report_data_no_records(config, test_db):
    """Verify report data structure and pattern resolution without .zip extension."""
    svc = ReportService(config, test_db)
    target_date = date(2026, 9, 4)
    data = svc.get_report_data(target_date)

    assert data["report_date"] == target_date
    assert data["batch_date_str"] == "September 04, 2026"
    assert data["checked_by"] == "Philip M. Bayudan"

    rows = data["table_rows"]
    assert len(rows) == 8
    # All rows should omit .zip
    for r in rows:
        assert not r["pattern"].endswith(".zip")
        assert not r["pattern"].endswith(".rar")
        assert r["verification_status"] == "Pending"
        assert r["integrity_check"] == "Not Applicable"

    # Row 1 TFS42PROD pattern resolution
    assert rows[0]["pattern"] == "TFS42PROD_20260904"
    assert rows[0]["file_type"] == "RAR"
    assert rows[0]["description"] == "Pre-Batch Clean Backup"


def test_report_data_with_completed_transfer(config, test_db):
    """Verify that completed transfer automates verification status and integrity check."""
    job = TransferJob(name="TFS42PROD", source_folder="/src", destination_folder="/dst")
    test_db.save_job(job)

    rec = TransferRecord(
        job_id=job.id,
        file_name="TFS42PROD_20260904.zip",
        source_path="/src/TFS42PROD_20260904.zip",
        destination_path="/dst/TFS42PROD_20260904.zip",
        file_size=10_485_760,  # 10 MB
        status=FileStatus.COMPLETED,
        transfer_completed=datetime(2026, 9, 4, 15, 30, 0),
        verification_passed=True,
    )
    test_db.save_record(rec)

    svc = ReportService(config, test_db)
    data = svc.get_report_data(date(2026, 9, 4))
    tfs_row = next(r for r in data["table_rows"] if r["job_name"] == "TFS42PROD")

    # Filename pattern must omit .zip extension
    assert tfs_row["pattern"] == "TFS42PROD_20260904"
    assert tfs_row["file_type"] == "RAR"
    assert tfs_row["file_size"] == "10.0 MB"
    assert tfs_row["verification_status"] == "Verified"
    assert tfs_row["integrity_check"] == "Passed"


def test_report_data_with_failed_transfer(config, test_db):
    """Verify that failed transfer sets status and integrity to Failed."""
    job = TransferJob(name="CSE", source_folder="/src", destination_folder="/dst")
    test_db.save_job(job)

    rec = TransferRecord(
        job_id=job.id,
        file_name="CSE_BACKUP_09-04-2026.zip",
        source_path="/src/CSE_BACKUP_09-04-2026.zip",
        destination_path="/dst/CSE_BACKUP_09-04-2026.zip",
        file_size=5_000_000,
        status=FileStatus.FAILED,
        transfer_completed=datetime(2026, 9, 4, 16, 0, 0),
        error_message="CRC mismatch detected",
        verification_passed=False,
    )
    test_db.save_record(rec)

    svc = ReportService(config, test_db)
    data = svc.get_report_data(date(2026, 9, 4))
    cse_row = next(r for r in data["table_rows"] if r["job_name"] == "CSE")

    assert cse_row["pattern"] == "CSE_BACKUP_09-04-2026"
    assert cse_row["verification_status"] == "Failed"
    assert cse_row["integrity_check"] == "Failed"
    assert "CRC mismatch" in cse_row["remarks"]


def test_generate_daily_report_excel(config, test_db, tmp_path):
    """Verify Excel workbook generation with proper header, table, control checks, and sign-off."""
    out_file = tmp_path / "test_report.xlsx"
    svc = ReportService(config, test_db)

    generated_path = svc.generate_daily_report(
        report_date=date(2026, 9, 4),
        operator_name="TEST_OP",
        checked_by="Philip M. Bayudan",
        repository_tag="TFSPH-PRIMARY-REPO",
        output_path=out_file,
    )

    assert generated_path.exists()
    wb = openpyxl.load_workbook(generated_path)
    ws = wb.active

    # Check Header Metadata
    assert ws["B5"].value == "September 04, 2026"
    assert ws["G5"].value == "TFSPH-PRIMARY-REPO"
    assert ws["B6"].value == "TEST_OP"

    # Check Table 1
    assert ws.cell(10, 1).value == 1
    assert ws.cell(10, 2).value == "TFS42PROD_20260904"
    assert ws.cell(10, 3).value == "RAR"
    assert ws.cell(17, 1).value == 8
    assert ws.cell(17, 2).value == "QMS_BACKUP_20260904"

    # Check Sign-Off
    assert ws.cell(32, 1).value == "TEST_OP"
    assert ws.cell(32, 4).value == "Philip M. Bayudan"


def test_dynamic_row_expansion_greater_than_8(config, test_db, tmp_path):
    """Verify dynamic row insertion when checklist has more than 8 systems."""
    systems = list(config.checklist_systems)
    systems.append({
        "no": 9,
        "job_name": "EXTRA1",
        "pattern": "EXTRA1_<YYYYMMDD>",
        "file_type": "ZIP",
        "description": "Extra 1 Backup",
        "expected_location": "",
    })
    systems.append({
        "no": 10,
        "job_name": "EXTRA2",
        "pattern": "EXTRA2_<YYYYMMDD>",
        "file_type": "ZIP",
        "description": "Extra 2 Backup",
        "expected_location": "",
    })
    config.checklist_systems = systems

    out_file = tmp_path / "expanded_report.xlsx"
    svc = ReportService(config, test_db)

    generated_path = svc.generate_daily_report(
        report_date=date(2026, 9, 4),
        operator_name="OPERATOR_EXP",
        output_path=out_file,
    )

    wb = openpyxl.load_workbook(generated_path)
    ws = wb.active

    # Table 1 should now have 10 rows (10 to 19)
    assert ws.cell(10, 1).value == 1
    assert ws.cell(18, 1).value == 9
    assert ws.cell(18, 2).value == "EXTRA1_20260904"
    assert ws.cell(19, 1).value == 10
    assert ws.cell(19, 2).value == "EXTRA2_20260904"

    # Daily Control Checks should have shifted by 2 rows (header at row 22)
    assert ws.cell(22, 1).value == "DAILY CONTROL CHECKS"
    # Check 1 is at row 24
    assert ws.cell(24, 1).value == 1
    assert "designated repository" in ws.cell(24, 2).value
    # Check 7 is at row 30
    assert ws.cell(30, 1).value == 7
    assert "escalated" in ws.cell(30, 2).value

    # Sign-Off should have shifted to row 34
    assert ws.cell(34, 1).value == "OPERATOR_EXP"
    assert ws.cell(34, 4).value == "Philip M. Bayudan"


def test_report_data_with_linked_job_mapping(config, test_db):
    """Verify that a test job (e.g. '001') can be mapped to a corporate system (e.g. 'TFS42PROD')."""
    # Create test job named 001
    job_001 = TransferJob(name="001", source_folder="/src/001", destination_folder="/dst/001")
    test_db.save_job(job_001)

    rec = TransferRecord(
        job_id=job_001.id,
        file_name="arbitrary_backup.exe",
        source_path="/src/001/arbitrary_backup.exe",
        destination_path="/dst/001/arbitrary_backup.exe",
        file_size=50_000_000,
        status=FileStatus.COMPLETED,
        transfer_completed=datetime(2026, 9, 4, 14, 0, 0),
        verification_passed=True,
    )
    test_db.save_record(rec)

    # Link row 1 (TFS42PROD) to test job 001
    systems = list(config.checklist_systems)
    systems[0]["linked_job"] = "001"
    config.checklist_systems = systems

    svc = ReportService(config, test_db)
    data = svc.get_report_data(date(2026, 9, 4))
    r1 = data["table_rows"][0]

    # Corporate name & pattern are preserved
    assert r1["job_name"] == "TFS42PROD"
    assert r1["pattern"] == "TFS42PROD_20260904"
    assert r1["file_type"] == "RAR"
    assert r1["description"] == "Pre-Batch Clean Backup"
    # Live stats pulled from test job 001
    assert r1["file_size"] == "47.7 MB"
    assert r1["completion_time"] == "02:00 PM"
    assert r1["verification_status"] == "Verified"
    assert r1["integrity_check"] == "Passed"
    assert r1["remarks"] == "Completed successfully"


def test_report_data_multi_file_aggregation(config, test_db):
    """Verify that multiple transferred files for a job aggregate correctly into total file size and latest time."""
    job = TransferJob(name="003", source_folder="/src/003", destination_folder="/dst/003")
    test_db.save_job(job)

    # 3 files: 10 MB, 20 MB, 30 MB (total 60 MB)
    rec1 = TransferRecord(
        job_id=job.id,
        file_name="file1.txt",
        source_path="/src/003/file1.txt",
        destination_path="/dst/003/file1.txt",
        file_size=10_485_760,  # 10 MB
        status=FileStatus.COMPLETED,
        transfer_completed=datetime(2026, 9, 4, 14, 0, 0),
        verification_passed=True,
    )
    rec2 = TransferRecord(
        job_id=job.id,
        file_name="file2.zip",
        source_path="/src/003/file2.zip",
        destination_path="/dst/003/file2.zip",
        file_size=20_971_520,  # 20 MB
        status=FileStatus.COMPLETED,
        transfer_completed=datetime(2026, 9, 4, 14, 15, 0),
        verification_passed=True,
    )
    rec3 = TransferRecord(
        job_id=job.id,
        file_name="file3.dat",
        source_path="/src/003/file3.dat",
        destination_path="/dst/003/file3.dat",
        file_size=31_457_280,  # 30 MB
        status=FileStatus.COMPLETED,
        transfer_completed=datetime(2026, 9, 4, 14, 30, 0),
        verification_passed=True,
    )
    test_db.save_record(rec1)
    test_db.save_record(rec2)
    test_db.save_record(rec3)

    # Link row 1 (TFS42PROD) to job 003
    systems = list(config.checklist_systems)
    systems[0]["linked_job"] = "003"
    config.checklist_systems = systems

    svc = ReportService(config, test_db)
    data = svc.get_report_data(date(2026, 9, 4))
    r = data["table_rows"][0]

    # Aggregated file size should be 60.0 MB
    assert r["file_size"] == "60.0 MB"
    # Latest completion time should be 02:30 PM
    assert r["completion_time"] == "02:30 PM"
    assert r["verification_status"] == "Verified"
    assert r["integrity_check"] == "Passed"
    assert "3 files verified" in r["remarks"]


